# -*- coding: utf-8 -*-
from __future__ import annotations

from datetime import date
from pathlib import Path
import unicodedata

import pandas as pd
from openpyxl import load_workbook


# =========================
# CONFIGURACIÓN
# =========================

INPUT_FILE = Path("data/merge_total.xlsx")
OUTPUT_FILE = Path("data/merge_total_con_mas_datos.xlsx")
REFERENCE_DATE = pd.Timestamp(date(2024, 12, 31))

DATE_COLUMNS = [
    "fecha_presentacion",
    "fecha_evaluacion",
    "fecha_nacimiento",
    "fecha_registro",
    "fecha_informacion",
    "fecha_email",
]


# =========================
# UTILIDADES
# =========================

def normalize_text(value) -> str | None:
    if pd.isna(value):
        return None
    text = str(value).strip()
    return text if text else None


def normalize_key(value) -> str | None:
    text = normalize_text(value)
    if text is None:
        return None
    text = unicodedata.normalize("NFD", text)
    text = "".join(ch for ch in text if unicodedata.category(ch) != "Mn")
    return text.lower().strip()


def calculate_age(fecha_nacimiento: pd.Timestamp | pd.NaT) -> int | None:
    if pd.isna(fecha_nacimiento):
        return None

    age = REFERENCE_DATE.year - fecha_nacimiento.year
    if (REFERENCE_DATE.month, REFERENCE_DATE.day) < (fecha_nacimiento.month, fecha_nacimiento.day):
        age -= 1
    return int(age)


def age_band(age: int | None) -> str | None:
    if age is None or pd.isna(age):
        return None

    if age <= 17:
        return "0-17"
    if age <= 24:
        return "18-24"
    if age <= 34:
        return "25-34"
    if age <= 44:
        return "35-44"
    if age <= 54:
        return "45-54"
    if age <= 64:
        return "55-64"
    return "65+"


def infer_estado_consolidado(row: pd.Series) -> str:
    estado = normalize_key(row["estado"])
    resultado = normalize_key(row["resultado_evaluacion"])
    estado_registro = normalize_key(row["estado_registro"])
    estado_informacion = normalize_key(row["estado_informacion"])
    estado_email = normalize_key(row["estado_email"])

    if estado == "pendiente":
        return "en evaluación"

    if estado == "evaluado" and resultado == "no_cumple":
        return "evaluado, no cumple"

    if estado == "evaluado" and resultado == "si_cumple" and estado_registro == "pendiente":
        return "evaluado, sí cumple, pendiente de registro"

    if (
        estado == "evaluado"
        and resultado == "si_cumple"
        and estado_registro == "registrado"
        and estado_informacion == "pendiente"
    ):
        return "registrado, pendiente de información"

    if (
        estado == "evaluado"
        and resultado == "si_cumple"
        and estado_registro == "registrado"
        and estado_informacion == "recibida"
        and estado_email == "pendiente"
    ):
        return "con información y pendiente de notificación"

    if (
        estado == "evaluado"
        and resultado == "si_cumple"
        and estado_registro == "registrado"
        and estado_informacion == "recibida"
        and estado_email == "enviado"
    ):
        return "notificado"

    return "estado no reconocido"


def calculate_fecha_fin_tramite(row: pd.Series) -> pd.Timestamp | pd.NaT:
    fecha_presentacion = row["fecha_presentacion"]
    fecha_evaluacion = row["fecha_evaluacion"]
    fecha_registro = row["fecha_registro"]
    fecha_informacion = row["fecha_informacion"]
    fecha_email = row["fecha_email"]

    estado = normalize_key(row["estado"])
    resultado = normalize_key(row["resultado_evaluacion"])

    if pd.isna(fecha_presentacion):
        return pd.NaT

    # Si no cumple, el trámite termina en evaluación
    if resultado == "no_cumple":
        return fecha_evaluacion if not pd.isna(fecha_evaluacion) else REFERENCE_DATE

    # Si sigue pendiente, cortar al "hoy ficticio"
    if estado == "pendiente":
        return REFERENCE_DATE

    # Si sí cumple pero no hay email, cortar al "hoy ficticio"
    if resultado == "si_cumple" and pd.isna(fecha_email):
        return REFERENCE_DATE

    # Caso general: usar la última fecha válida del flujo
    for candidate in [fecha_email, fecha_informacion, fecha_registro, fecha_evaluacion]:
        if not pd.isna(candidate):
            return candidate

    return REFERENCE_DATE


def calculate_dias_tramite(row: pd.Series) -> int | None:
    fecha_presentacion = row["fecha_presentacion"]
    fecha_fin_tramite = row["fecha_fin_tramite"]

    if pd.isna(fecha_presentacion) or pd.isna(fecha_fin_tramite):
        return None

    return int((fecha_fin_tramite - fecha_presentacion).days)


def dias_tramite_band(dias: int | None) -> str | None:
    if dias is None or pd.isna(dias):
        return None

    if dias < 0:
        return "valor inconsistente"
    if dias <= 7:
        return "0-7"
    if dias <= 15:
        return "8-15"
    if dias <= 30:
        return "16-30"
    if dias <= 60:
        return "31-60"
    if dias <= 90:
        return "61-90"
    if dias <= 180:
        return "91-180"
    return "181+"


def reorder_columns(df: pd.DataFrame) -> pd.DataFrame:
    # Bloque solicitante a la izquierda
    solicitante_cols = [
        "codigo_solicitante",
        "nombre",
        "apellido",
        "sexo",
        "fecha_nacimiento",
        "edad",
        "rango_edad",
        "nivel_de_estudios",
        "ocupacion",
    ]

    # Bloque solicitud y trámite a la derecha
    tramite_cols = [
        "codigo_solicitud",
        "fecha_presentacion",
        "estado",
        "fecha_evaluacion",
        "resultado_evaluacion",
        "estado_consolidado",
        "estado_registro",
        "fecha_registro",
        "estado_informacion",
        "fecha_informacion",
        "estado_email",
        "fecha_email",
        "fecha_fin_tramite",
        "dias_tramite",
        "categoria_dias_tramite",
    ]

    known_order = solicitante_cols + tramite_cols
    existing_known = [col for col in known_order if col in df.columns]
    remaining = [col for col in df.columns if col not in existing_known]

    return df[existing_known + remaining]


def apply_excel_formatting(output_file: Path) -> None:
    wb = load_workbook(output_file)
    ws = wb[wb.sheetnames[0]]

    date_headers = {
        "fecha_nacimiento",
        "fecha_presentacion",
        "fecha_evaluacion",
        "fecha_registro",
        "fecha_informacion",
        "fecha_email",
        "fecha_fin_tramite",
    }

    header_map = {ws.cell(row=1, column=i).value: i for i in range(1, ws.max_column + 1)}

    for header in date_headers:
        col_idx = header_map.get(header)
        if not col_idx:
            continue
        for row in range(2, ws.max_row + 1):
            ws.cell(row=row, column=col_idx).number_format = "DD/MM/YYYY"

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    wb.save(output_file)


# =========================
# PROCESO PRINCIPAL
# =========================

def main() -> None:
    if not INPUT_FILE.exists():
        raise FileNotFoundError(f"No se encontró el archivo de entrada: {INPUT_FILE}")

    df = pd.read_excel(INPUT_FILE)

    expected_columns = [
        "codigo_solicitud",
        "codigo_solicitante",
        "fecha_presentacion",
        "estado",
        "fecha_evaluacion",
        "resultado_evaluacion",
        "nombre",
        "apellido",
        "sexo",
        "fecha_nacimiento",
        "nivel_de_estudios",
        "ocupacion",
        "estado_registro",
        "fecha_registro",
        "estado_informacion",
        "fecha_informacion",
        "estado_email",
        "fecha_email",
    ]

    missing = [col for col in expected_columns if col not in df.columns]
    if missing:
        raise ValueError(f"Faltan columnas esperadas en el archivo: {missing}")

    # Convertir fechas
    for col in DATE_COLUMNS:
        df[col] = pd.to_datetime(df[col], errors="coerce")

    # Nuevas columnas
    df["edad"] = df["fecha_nacimiento"].apply(calculate_age)
    df["rango_edad"] = df["edad"].apply(age_band)
    df["estado_consolidado"] = df.apply(infer_estado_consolidado, axis=1)
    df["fecha_fin_tramite"] = df.apply(calculate_fecha_fin_tramite, axis=1)
    df["dias_tramite"] = df.apply(calculate_dias_tramite, axis=1)
    df["categoria_dias_tramite"] = df["dias_tramite"].apply(dias_tramite_band)

    # Reordenar columnas
    df = reorder_columns(df)

    # Crear carpeta de salida si no existe
    OUTPUT_FILE.parent.mkdir(parents=True, exist_ok=True)

    # Guardar Excel
    df.to_excel(OUTPUT_FILE, index=False)
    apply_excel_formatting(OUTPUT_FILE)

    print(f"Archivo generado correctamente en: {OUTPUT_FILE}")


if __name__ == "__main__":
    main()